from pathlib import Path
import sys,json,zipfile,xml.etree.ElementTree as ET,re
import openpyxl
ROOT=Path('/mnt/BioDisk/wangmingyuan/GI_multiomics_screening')
for p in ROOT.glob('*/00_docs/*full.xml'):
 try:
  t=ET.parse(p); text='\n'.join(''.join(x.itertext()) for x in t.findall('.//p'))
  p.with_suffix('.txt').write_text(text)
  print('\nARTICLE',p.name)
  for para in text.splitlines():
   if re.search(r'normaliz|imput|missing|TMT|localiz|Data availability|censor|survival|spectronaut|DIA-NN',para,re.I):print(para)
  print('ATTACHMENTS',[(x.attrib,''.join(x.itertext())[:180]) for x in t.findall('.//supplementary-material')])
 except Exception as e:print(e)
for project in ['A_CRLM_phosphoregulation','B_ICC_lymphatic_invasion']:
 for p in (ROOT/project/'02_raw').rglob('*'):
  if p.suffix not in ['.xlsx','.xls']:continue
  print('\nWORKBOOK',p)
  try:
   z=zipfile.ZipFile(p)
   if 'xl/workbook.bin' in z.namelist():
    from pyxlsb import open_workbook
    with open_workbook(str(p)) as w:
     print('XLSB',w.sheets)
     for s in w.sheets:
      with w.get_sheet(s) as sh:
       rr=[];n=0
       for row in sh.rows():
        if n<5:rr.append([c.v for c in row][:18])
        n+=1
       print(s,n,rr)
   else:
    with p.open('rb') as f:
     w=openpyxl.load_workbook(f,read_only=True,data_only=True)
     for s in w:
      declared=(s.max_row,s.max_column);s.reset_dimensions();rr=[];n=0;nc=0
      for row in s.iter_rows(values_only=True):
       if n<5:rr.append(list(row)[:25])
       n+=1;nc=max(nc,len(row))
      print(s.title,'declared',declared,'actual',(n,nc),'rows',rr)
  except Exception as e:print(type(e).__name__,e)
print('SOURCE_STRUCTURE_AUDIT_COMPLETE')
